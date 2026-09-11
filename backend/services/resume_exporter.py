from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


# --------------------------------------------------
# PDF Export
# --------------------------------------------------

def generate_resume_pdf(resume: dict) -> BytesIO:
    buffer = BytesIO()

    pdf = canvas.Canvas(
        buffer,
        pagesize=A4
    )

    width, height = A4

    x = 50
    y = height - 50

    def write_line(
        text,
        font="Helvetica",
        size=11,
        gap=16
    ):
        nonlocal y

        if y < 60:
            pdf.showPage()
            y = height - 50

        pdf.setFont(font, size)
        pdf.drawString(x, y, str(text))
        y -= gap

    # Title
    write_line(
        resume.get(
            "target_job_title",
            "Tailored Resume"
        ),
        font="Helvetica-Bold",
        size=18,
        gap=28
    )

    # Summary
    write_line(
        "Professional Summary",
        font="Helvetica-Bold",
        size=13,
        gap=20
    )

    summary = resume.get(
        "professional_summary",
        ""
    )

    for line in wrap_text(summary, 90):
        write_line(line)

    y -= 8

    # Skills
    write_line(
        "Skills",
        font="Helvetica-Bold",
        size=13,
        gap=20
    )

    skills = resume.get(
        "skills",
        []
    )

    if skills:
        skills_text = ", ".join(skills)

        for line in wrap_text(
            skills_text,
            90
        ):
            write_line(line)

    y -= 8

    # Projects
    write_line(
        "Projects",
        font="Helvetica-Bold",
        size=13,
        gap=20
    )

    for project in resume.get(
        "projects",
        []
    ):
        write_line(
            project.get(
                "name",
                "Project"
            ),
            font="Helvetica-Bold",
            size=11
        )

        technologies = project.get(
            "technologies",
            []
        )

        if technologies:
            tech_text = (
                "Technologies: "
                + ", ".join(
                    technologies
                )
            )

            for line in wrap_text(
                tech_text,
                90
            ):
                write_line(line)

        for bullet in project.get(
            "bullet_points",
            []
        ):
            for line in wrap_text(
                f"- {bullet}",
                85
            ):
                write_line(line)

        y -= 6

    # Education
    write_line(
        "Education",
        font="Helvetica-Bold",
        size=13,
        gap=20
    )

    for item in resume.get(
        "education",
        []
    ):
        if isinstance(item, dict):
            degree = item.get(
                "degree",
                ""
            )

            university = item.get(
                "university",
                ""
            )

            text = degree

            if university:
                text += (
                    f" - {university}"
                )

        else:
            text = str(item)

        write_line(text)

    y -= 8

    # Certifications
    write_line(
        "Certifications",
        font="Helvetica-Bold",
        size=13,
        gap=20
    )

    for cert in resume.get(
        "certifications",
        []
    ):
        write_line(
            f"- {cert}"
        )

    pdf.save()

    buffer.seek(0)

    return buffer


# --------------------------------------------------
# DOCX Export
# --------------------------------------------------

def generate_resume_docx(
    resume: dict
) -> BytesIO:
    buffer = BytesIO()

    document = Document()

    document.add_heading(
        resume.get(
            "target_job_title",
            "Tailored Resume"
        ),
        level=1
    )

    # Summary
    document.add_heading(
        "Professional Summary",
        level=2
    )

    document.add_paragraph(
        resume.get(
            "professional_summary",
            ""
        )
    )

    # Skills
    document.add_heading(
        "Skills",
        level=2
    )

    skills = resume.get(
        "skills",
        []
    )

    document.add_paragraph(
        ", ".join(skills)
    )

    # Projects
    document.add_heading(
        "Projects",
        level=2
    )

    for project in resume.get(
        "projects",
        []
    ):
        document.add_heading(
            project.get(
                "name",
                "Project"
            ),
            level=3
        )

        technologies = (
            project.get(
                "technologies",
                []
            )
        )

        if technologies:
            document.add_paragraph(
                "Technologies: "
                + ", ".join(
                    technologies
                )
            )

        for bullet in project.get(
            "bullet_points",
            []
        ):
            document.add_paragraph(
                bullet,
                style="List Bullet"
            )

    # Education
    document.add_heading(
        "Education",
        level=2
    )

    for item in resume.get(
        "education",
        []
    ):
        if isinstance(item, dict):
            degree = item.get(
                "degree",
                ""
            )

            university = item.get(
                "university",
                ""
            )

            text = degree

            if university:
                text += (
                    f" - {university}"
                )

        else:
            text = str(item)

        document.add_paragraph(
            text
        )

    # Certifications
    document.add_heading(
        "Certifications",
        level=2
    )

    for cert in resume.get(
        "certifications",
        []
    ):
        document.add_paragraph(
            cert,
            style="List Bullet"
        )

    document.save(buffer)

    buffer.seek(0)

    return buffer


# --------------------------------------------------
# Helper
# --------------------------------------------------

def wrap_text(
    text: str,
    max_chars: int
):
    words = str(text).split()

    lines = []
    current_line = []

    current_length = 0

    for word in words:
        word_length = len(word)

        if (
            current_length
            + word_length
            + 1
            > max_chars
        ):
            lines.append(
                " ".join(
                    current_line
                )
            )

            current_line = [
                word
            ]

            current_length = (
                word_length
            )

        else:
            current_line.append(
                word
            )

            current_length += (
                word_length + 1
            )

    if current_line:
        lines.append(
            " ".join(
                current_line
            )
        )

    return lines